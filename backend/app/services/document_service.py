import os
import re
import json
import logging
import requests
from typing import List, Dict, Any, Optional

logger = logging.getLogger("jam_analyzer")

# Conditional imports for document parsers
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    import pptx
except ImportError:
    pptx = None

from app.config import settings
from app.utils.ai_parser import (
    parse_ai_response,
    DocumentAnalysisModel,
    TopicsListModel,
    SpeechToDocumentEvaluationModel,
    VivaQuestionsModel,
    VivaEvaluationModel
)

class DocumentService:
    def __init__(self):
        # All AI calls route through GroqService via _call_llm_json.
        # No additional client initialization is required here.
        pass

    def extract_text(self, file_path: str, filename: str) -> str:
        """
        Extracts raw text content from the uploaded file based on its extension.
        Logs detailed extraction status.
        """
        ext = os.path.splitext(filename.lower())[1]
        text_content = ""

        # Step 1: Diagnose input file details
        try:
            file_size_bytes = os.path.getsize(file_path)
            file_size_kb = round(file_size_bytes / 1024, 2)
        except Exception:
            file_size_kb = 0.0

        logger.info("--- Document Extraction Diagnostics ---")
        logger.info(f"File Name: {filename}")
        logger.info(f"File Type: {ext}")
        logger.info(f"File Size: {file_size_kb} KB")
        logger.info(f"Storage Path: {file_path}")

        try:
            if ext == ".pdf":
                if pypdf:
                    reader = pypdf.PdfReader(file_path)
                    text_content = "\n".join([page.extract_text() or "" for page in reader.pages])
                else:
                    logger.warning("pypdf is not installed. Using basic byte decoding fallback.")
                    text_content = self._fallback_text_extraction(file_path)
            
            elif ext in [".docx", ".doc"]:
                if docx and ext == ".docx":
                    doc = docx.Document(file_path)
                    text_content = "\n".join([p.text for p in doc.paragraphs])
                else:
                    logger.warning("docx parser fallback triggered.")
                    text_content = self._fallback_text_extraction(file_path)
            
            elif ext in [".pptx", ".ppt"]:
                if pptx and ext == ".pptx":
                    prs = pptx.Presentation(file_path)
                    slides_text = []
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text") and shape.text:
                                slides_text.append(shape.text)
                    text_content = "\n".join(slides_text)
                else:
                    logger.warning("pptx parser fallback triggered.")
                    text_content = self._fallback_text_extraction(file_path)
            
            elif ext in [".txt", ".md"]:
                # Try standard UTF-8 first
                try:
                    with open(file_path, "r", encoding="utf-8") as f:
                        text_content = f.read()
                except UnicodeDecodeError:
                    # Try Latin-1 fallback
                    logger.warning(f"UTF-8 decoding failed for {filename}. Trying Latin-1 fallback.")
                    try:
                        with open(file_path, "r", encoding="latin-1") as f:
                            text_content = f.read()
                    except Exception as e:
                        raise ValueError(f"Unsupported Encoding: Failed to read text file. Details: {e}")
            
            else:
                # Catch-all basic reading
                text_content = self._fallback_text_extraction(file_path)

        except Exception as e:
            logger.error(f"Error during file text extraction: {e}")
            raise ValueError(f"Text Extraction Failed: {str(e)}")

        # Cleanup text formatting slightly
        text_content = re.sub(r'\s+', ' ', text_content).strip()
        
        # Step 2: Validate extracted character count
        text_len = len(text_content)
        if text_len == 0:
            raise ValueError(f"No Content Extracted: The file '{filename}' was resolved as empty.")

        preview = text_content[:200] + "..." if text_len > 200 else text_content
        logger.info(f"Extracted Characters: {text_len}")
        logger.info(f"Extracted Text Preview: \"{preview}\"")
        logger.info("---------------------------------------")
            
        return text_content

    def _fallback_text_extraction(self, file_path: str) -> str:
        """
        Reads binary files extracting printable ASCII strings as fallback.
        """
        try:
            with open(file_path, "rb") as f:
                data = f.read(500000) # Limit to first 500KB for safety
                ascii_strings = re.findall(b"[a-zA-Z0-9\s\.,;:!@#\$%\^&\*\(\)\-_\+=\[\]\{\}<>\?\/\\|~`']{4,}", data)
                decoded = []
                for s in ascii_strings:
                    try:
                        decoded.append(s.decode("utf-8"))
                    except:
                        pass
                return "\n".join(decoded)
        except Exception as e:
            logger.error(f"Fallback text extraction failed: {e}")
            return ""

    def _call_llm_json(
        self, prompt: str, system_prompt: str = "You are an elite coach. Output clean JSON only.", model_class: Any = None
    ) -> Dict[str, Any]:
        """
        Routes LLM requests to the centralized GroqService via text call,
        logs raw response, cleans/parses/normalizes, and retries once if json loading fails.
        """
        from app.services.groq_service import GroqService
        from app.utils.ai_parser import clean_raw_string
        
        groq_service = GroqService()
        
        logger.info("[_call_llm_json] Sending request to Groq...")
        
        # Attempt 1
        raw_response = ""
        try:
            raw_response = groq_service.call_groq_text(prompt, system_prompt)
            # Try to test JSON loading immediately to determine if retry is needed
            json.loads(clean_raw_string(raw_response))
            # Run the parser
            parsed = parse_ai_response(raw_response, model_class)
            return parsed
        except Exception as e:
            logger.warning(f"[_call_llm_json] First attempt failed JSON check: {e}. Retrying once...")
            
        # Retry Attempt (Requirement 7)
        logger.info("[_call_llm_json] Retry attempts: 1")
        try:
            raw_response = groq_service.call_groq_text(prompt, system_prompt)
            parsed = parse_ai_response(raw_response, model_class)
            return parsed
        except Exception as e:
            logger.error(f"[_call_llm_json] Retry attempt failed: {e}")
            # Construct absolute fallback dictionary from target model
            fallback = model_class.model_validate({})
            return fallback.model_dump()

    def analyze_document_content(self, text: str) -> Dict[str, Any]:
        """
        AI reads the extracted text and outputs title, summary, key_concepts, keywords, and classification.
        """
        if not text.strip():
            raise ValueError("Empty Document: Content to analyze is empty.")

        truncated_text = text[:15000] # Limit tokens sent for summary
        
        prompt = f"""
        Analyze the following extracted document text and return a JSON object.
        
        Extracted Document Text:
        \"\"\"{truncated_text}\"\"\"

        Return ONLY a valid JSON object.
        Do not include Markdown.
        Do not include explanations.
        Do not include code fences.
        Return exactly this schema:
        {{
          "title": "Descriptive title of the document",
          "summary": "Concise 2-3 sentence summary of the document",
          "key_concepts": ["Concept 1", "Concept 2", "Concept 3"],
          "keywords": ["Keyword 1", "Keyword 2"],
          "learning_objectives": ["Objective 1", "Objective 2"],
          "is_project_report": false,
          "is_resume": false
        }}
        """

        try:
            return self._call_llm_json(
                prompt, 
                "You are a professional document parser and analyst. Output clean JSON only.",
                DocumentAnalysisModel
            )
        except Exception as e:
            logger.error(f"Failed document analysis: {e}")
            raise ValueError(f"AI Request Failed: Failed to analyze document content. Details: {str(e)}")

    def generate_topics(self, doc_analysis: Dict[str, Any], text: str) -> List[Dict[str, Any]]:
        """
        Generates 5-10 unique speaking topics from the document based on its category and contents.
        """
        title = doc_analysis.get("title", "Document")
        summary = doc_analysis.get("summary", "")
        concepts = ", ".join(doc_analysis.get("key_concepts", []))
        keywords = ", ".join(doc_analysis.get("keywords", []))

        # Determine type
        doc_type = "standard"
        if doc_analysis.get("is_project_report"):
            doc_type = "project_report"
        elif doc_analysis.get("is_resume"):
            doc_type = "resume"

        prompt = f"""
        You are a topic generation agent for Document-Based Presentation training.
        Based on this document:
        Title: {title}
        Summary: {summary}
        Key Concepts: {concepts}
        Keywords: {keywords}
        Document Classification: {doc_type.upper()}
        Extracted content preview:
        \"\"\"{text[:4000]}\"\"\"

        Generate a list of 5 to 10 unique, relevant speaking presentation topics strictly based on the extracted content.
        Ensure you match the Document Classification requirements:
        - If STANDARD: Generate conceptual topics mapping the document's themes. (Beginner, Intermediate, Advanced).
        - If PROJECT_REPORT: Automatically generate specific topics:
          1. Project Presentation Topic
          2. Architecture Design and Choices
          3. Technical Implementation Details
          4. Challenges & Future Improvements
        - If RESUME: Automatically generate professional topics:
          1. Professional Experience & Role Highlights
          2. Key Projects & Technical Achievements
          3. Career Goals & Fit
          4. Behavioral strengths based on resume

        Return ONLY a valid JSON object.
        Do not include Markdown.
        Do not include explanations.
        Do not include code fences.
        Return exactly this schema:
        {{
          "topics": [
            {{
              "topic": "The speaking topic title",
              "category": "Technology/Engineering/Career/etc.",
              "difficulty": "Beginner/Intermediate/Advanced",
              "talking_points": ["point 1", "point 2", "point 3"],
              "estimated_speaking_time": 60,
              "keywords": ["key1", "key2", "key3"]
            }}
          ]
        }}
        """

        try:
            res_data = self._call_llm_json(
                prompt, 
                "You are a communication training coordinator. Output clean JSON only.",
                TopicsListModel
            )
            
            topics = res_data.get("topics", [])
            if not topics:
                raise ValueError("AI returned empty topics array.")
            return topics
        except Exception as e:
            logger.error(f"Failed to generate topics: {e}")
            raise ValueError(f"AI Request Failed: Failed to generate topics. Details: {str(e)}")

    def evaluate_speech_against_document(
        self, document_text: str, topic_title: str, transcript: str, communication_scores: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        Performs comparative AI analysis between document content and user speaking transcript.
        Evaluates Content Accuracy, Concept Coverage, Topic Relevance, Understanding, Explanation Quality, and Tech Correctness.
        """
        truncated_doc = document_text[:12000] # Limit token inputs
        
        prompt = f"""
        You are an elite academic presentation evaluator and document presentation coach.
        Compare the user's spoken transcript against the source document to evaluate their content understanding.

        Topic User Chose to Present: "{topic_title}"
        User spoken transcript:
        \"\"\"{transcript}\"\"\"

        Source Document text:
        \"\"\"{truncated_doc}\"\"\"

        Return ONLY a valid JSON object.
        Do not include Markdown.
        Do not include explanations.
        Do not include code fences.
        Return exactly this schema:
        {{
          "accuracy_score": 75,
          "coverage_score": 70,
          "understanding_score": 72,
          "explanation_quality": 74,
          "technical_correctness": 76,
          "relevance_score": 80,
          "knowledge_gaps": [
            {{
              "concept": "Name of the missing concept/topic",
              "description": "1 sentence describing what was missed and where to find it"
            }}
          ],
          "suggested_improvements": ["improvement 1", "improvement 2"],
          "coach_recommendations": ["recommendation 1", "recommendation 2"],
          "follow_up_questions": ["question 1", "question 2", "question 3"]
        }}
        """

        try:
            return self._call_llm_json(
                prompt, 
                "You are a master presentation examiner. Output clean JSON only.",
                SpeechToDocumentEvaluationModel
            )
        except Exception as e:
            logger.error(f"Failed comparative speech-to-document evaluation: {e}")
            raise ValueError(f"AI Request Failed: Failed to compare speech against document. Details: {str(e)}")

    def generate_viva_questions(self, document_text: str, mode: str) -> List[str]:
        """
        Generates 5 oral examination / interview questions based on document text and chosen mode.
        """
        truncated_doc = document_text[:12000]
        
        prompt = f"""
        You are an expert interviewer and academic coordinator.
        Based on the following document content:
        \"\"\"{truncated_doc}\"\"\"

        Generate exactly 5 challenging oral questions matching the mode: {mode.upper()}
        - VIVA: Focus on theory, validity, methodology, academic defense.
        - PROJECT: Focus on system design, database schemas, code flow, scaling, testing.
        - RESUME: Focus on experience highlights, candidate skills verification, technical questions about projects, HR behavioral fit.

        Return ONLY a valid JSON object.
        Do not include Markdown.
        Do not include explanations.
        Do not include code fences.
        Return exactly this schema:
        {{
          "questions": [
            "Question 1?",
            "Question 2?",
            "Question 3?",
            "Question 4?",
            "Question 5?"
          ]
        }}
        """

        try:
            res_data = self._call_llm_json(
                prompt, 
                "You are a professional panel examiner. Output clean JSON only.",
                VivaQuestionsModel
            )
            questions = res_data.get("questions", [])
            if not questions or len(questions) < 5:
                raise ValueError(f"AI returned invalid questions list: {questions}")
            return questions
        except Exception as e:
            logger.error(f"Failed to generate viva questions: {e}")
            raise ValueError(f"AI Request Failed: Failed to generate viva questions. Details: {str(e)}")

    def evaluate_viva_response(self, document_text: str, question: str, answer_transcript: str) -> Dict[str, Any]:
        """
        Grades a single viva oral answer out of 100 and provides specific qualitative feedback.
        """
        truncated_doc = document_text[:12000]
        
        prompt = f"""
        Evaluate this candidate's oral response to a Viva question based on the source document:
        
        Question: "{question}"
        Candidate Spoken Response:
        \"\"\"{answer_transcript}\"\"\"

        Source Document:
        \"\"\"{truncated_doc}\"\"\"

        Return ONLY a valid JSON object.
        Do not include Markdown.
        Do not include explanations.
        Do not include code fences.
        Return exactly this schema:
        {{
          "score": 75,
          "feedback": "2-sentence summary feedback",
          "correct_elements": ["element 1", "element 2"],
          "incorrect_elements": ["weakness 1"],
          "ideal_response": "Sample expert-level response to the question"
        }}
        """

        try:
            return self._call_llm_json(
                prompt, 
                "You are an oral board assessor. Output clean JSON only.",
                VivaEvaluationModel
            )
        except Exception as e:
            logger.error(f"Failed to evaluate viva response: {e}")
            raise ValueError(f"AI Request Failed: Failed to grade viva answer. Details: {str(e)}")
